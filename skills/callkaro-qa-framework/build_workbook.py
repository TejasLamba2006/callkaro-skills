#!/usr/bin/env python3
"""Build the standard CallKaro QA workbook from one JSON file.

    python build_workbook.py qa.json out.xlsx

JSON keys (all optional except summary): see WORKBOOK.md. Scores are live formulas,
so testers editing Status / passed / score cells update the Summary tab.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HEAD = PatternFill("solid", fgColor="1F2937")
BAND = PatternFill("solid", fgColor="F3F4F6")
GREEN, RED, AMBER, GREY = "C6EFCE", "FFC7CE", "FFEB9C", "E5E7EB"
CATS = ["DBEAFE", "FCE7F3", "FEF3C7", "D1FAE5", "EDE9FE", "FFEDD5", "CFFAFE", "FEE2E2",
        "ECFCCB", "E0E7FF", "FAE8FF", "F5F5F4"]
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUSES = ["Pass", "Fail", "Blocked", "Not Run"]
LAYERS = ["prompt", "function", "config", "voice/platform", "data", "none"]

# (header, json key, width); "pass %" is computed, not read
REVIEW = [("Area", "area", 18), ("Check", "check", 34), ("Score (0-10)", "score", 11),
          ("What is good", "good", 40), ("What is bad", "bad", 40),
          ("Recommendation", "recommendation", 40), ("Reference", "reference", 22)]
SIM = [("Test id", "id", 12), ("Test case name", "name", 44), ("Category", "category", 18),
       ("Persona", "persona", 22), ("Branch / situation", "branch", 36), ("Critical", "critical", 9),
       ("Runs", "runs", 7), ("Passed", "passed", 8), ("Pass %", None, 9),
       ("Previous pass %", "previous_pct", 11), ("What happened", "happened", 44),
       ("Comment on agent handling", "comment", 36), ("Layer / fault", "layer", 14),
       ("Batch id", "batch_id", 14)]
MANUAL = [("TC id", "id", 8), ("Category", "category", 18), ("Persona / setup", "persona", 30),
          ("Branch / situation", "branch", 34), ("What to say", "say", 44),
          ("Expected agent behaviour (incl. how the call ends)", "expected", 52),
          ("Critical", "critical", 9), ("Status", "status", 10), ("Score (0-10)", "score", 11),
          ("Comments", "comments", 36), ("Layer / fault", "layer", 14), ("Call id", "call_id", 14)]


def sheet(wb, title, cols, rows, freeze_col="A"):
    ws = wb.create_sheet(title)
    for c, (h, _, w) in enumerate(cols, 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font, cell.border = HEAD, Font(bold=True, color="FFFFFF"), BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 32
    for r, row in enumerate(rows, 2):
        for c, (h, key, _) in enumerate(cols, 1):
            v = row.get(key) if key else None
            cell = ws.cell(r, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BOX
            if r % 2 == 1:
                cell.fill = BAND
    ws.freeze_panes = f"{chr(ord(freeze_col) + 1)}2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(cols)).column_letter}{max(len(rows) + 1, 2)}"
    return ws


def col(cols, header):
    return chr(ord("A") + [h for h, _, _ in cols].index(header))


def dropdown(ws, letter, values, n):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{n + 60}")  # spare rows for cases added later


def paint_cats(ws, letter, n, colors):
    for r in range(2, n + 2):
        cell = ws[f"{letter}{r}"]
        if cell.value in colors:
            cell.fill = PatternFill("solid", fgColor=colors[cell.value])
            cell.font = Font(bold=True)


def common_rules(ws, cols, n, crit=True, status=False):
    end = n + 60
    if crit:
        L = col(cols, "Critical")
        ws.conditional_formatting.add(f"{L}2:{L}{end}", CellIsRule(
            operator="equal", formula=['"Y"'], font=Font(bold=True, color="9C0006"),
            fill=PatternFill("solid", bgColor=RED)))
        dropdown(ws, L, ["Y", "N"], n)
    if status:
        L = col(cols, "Status")
        for val, color in zip(STATUSES, [GREEN, RED, AMBER, GREY]):
            ws.conditional_formatting.add(f"{L}2:{L}{end}", CellIsRule(
                operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", bgColor=color)))
        dropdown(ws, L, STATUSES, n)
    L = col(cols, "Layer / fault") if any(h == "Layer / fault" for h, _, _ in cols) else None
    if L:
        dropdown(ws, L, LAYERS, n)


def scale(ws, rng, lo, hi):
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B",
        mid_type="num", mid_value=lo, mid_color="FFEB84",
        end_type="num", end_value=hi, end_color="63BE7B"))


def build(data, out):
    wb = Workbook()
    wb.remove(wb.active)
    s = data["summary"]
    sim, man, rev = data.get("sim", []), data.get("manual", []), data.get("review", [])

    cats, colors = [], {}
    for row in sim + man:
        c = row.get("category")
        if c and c not in colors:
            colors[c] = CATS[len(colors) % len(CATS)]
            cats.append(c)

    summ = wb.create_sheet("Summary")

    # Agent Brief
    ws = sheet(wb, "Agent Brief", [("Section", "section", 24), ("Detail", "detail", 110)],
               data.get("brief", []))
    for r in range(2, len(data.get("brief", [])) + 2):
        ws[f"A{r}"].font = Font(bold=True)

    # Agent Review
    ws = sheet(wb, "Agent Review", REVIEW, rev)
    scale(ws, f"C2:C{len(rev) + 60}", 6, 8.5)

    # Simulation Results
    ws = sheet(wb, "Simulation Results", SIM, sim, "B")
    P, R, Q = col(SIM, "Pass %"), col(SIM, "Runs"), col(SIM, "Passed")
    for r in range(2, len(sim) + 2):
        ws[f"{P}{r}"] = f'=IF(N({R}{r})>0,{Q}{r}/{R}{r},"")'
        ws[f"{P}{r}"].number_format = "0%"
        ws[f"{col(SIM, 'Previous pass %')}{r}"].number_format = "0%"
    scale(ws, f"{P}2:{P}{len(sim) + 60}", 0.6, 0.85)
    paint_cats(ws, col(SIM, "Category"), len(sim), colors)
    common_rules(ws, SIM, len(sim))
    sim_n = len(sim)

    # Manual
    ws = sheet(wb, "Manual Test Cases", MANUAL, man)
    for r in range(2, len(man) + 2):
        if ws[f"{col(MANUAL, 'Status')}{r}"].value is None:
            ws[f"{col(MANUAL, 'Status')}{r}"] = "Not Run"
    scale(ws, f"{col(MANUAL, 'Score (0-10)')}2:{col(MANUAL, 'Score (0-10)')}{len(man) + 60}", 6, 8.5)
    paint_cats(ws, col(MANUAL, "Category"), len(man), colors)
    common_rules(ws, MANUAL, len(man), status=True)

    # Summary (live formulas)
    sim_end, rev_end = sim_n + 60, len(rev) + 60
    cr = f"'Simulation Results'!{col(SIM, 'Critical')}2:{col(SIM, 'Critical')}{sim_end}"
    pc = f"'Simulation Results'!{P}2:{P}{sim_end}"
    rows = [
        ("Agent", s.get("agent_name")), ("Agent id", s.get("agent_id")),
        ("Version under QA", f"{s.get('version_name', '')} ({s.get('version_id', '')})"),
        ("Date", s.get("date")),
        ("Review %", f"=IFERROR(AVERAGE('Agent Review'!C2:C{rev_end})*10,0)"),
        ("Simulation %", f"=IFERROR(AVERAGE({pc})*100,0)"),
        ("QA score", "=0.5*B6+0.5*B7"),
        ("Critical cases at 5/5 (all runs passed)",
         f'=COUNTIFS({cr},"Y",{pc},1)&" / "&COUNTIF({cr},"Y")'),
        ("Verdict", f'=IF(AND(B8>=85,COUNTIFS({cr},"Y",{pc},"<1")=0),"PASS","FAIL")'),
        ("Manual: passed / run",
         f"=COUNTIF('Manual Test Cases'!{col(MANUAL, 'Status')}2:{col(MANUAL, 'Status')}{len(man) + 60},\"Pass\")"
         f"&\" / \"&(COUNTA('Manual Test Cases'!A2:A{len(man) + 60})"
         f"-COUNTIF('Manual Test Cases'!{col(MANUAL, 'Status')}2:{col(MANUAL, 'Status')}{len(man) + 60},\"Not Run\"))"),
    ]
    summ["A1"], summ["A1"].font = "CallKaro Voice Agent QA", Font(bold=True, size=16)
    for i, (k, v) in enumerate(rows, 2):
        summ[f"A{i}"], summ[f"B{i}"] = k, v
        summ[f"A{i}"].font = Font(bold=True)
    for r in (6, 7, 8):
        summ[f"B{r}"].number_format = "0.0"
    summ.conditional_formatting.add("B10", CellIsRule(operator="equal", formula=['"PASS"'],
                                    fill=PatternFill("solid", bgColor="63BE7B"), font=Font(bold=True)))
    summ.conditional_formatting.add("B10", CellIsRule(operator="equal", formula=['"FAIL"'],
                                    fill=PatternFill("solid", bgColor="F8696B"), font=Font(bold=True)))
    for r in (6, 7, 8):
        summ.conditional_formatting.add(f"B{r}", ColorScaleRule(
            start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=60,
            mid_color="FFEB84", end_type="num", end_value=85, end_color="63BE7B"))
    summ["A13"], summ["A13"].font = "Cases by category (legend)", Font(bold=True, size=12)
    for c, h in enumerate(["Category", "Simulation", "Manual"], 1):
        cell = summ.cell(14, c, h)
        cell.fill, cell.font = HEAD, Font(bold=True, color="FFFFFF")
    sc, mc = col(SIM, "Category"), col(MANUAL, "Category")
    for i, cat in enumerate(cats, 15):
        summ[f"A{i}"] = cat
        summ[f"A{i}"].fill = PatternFill("solid", fgColor=colors[cat])
        summ[f"B{i}"] = f"=COUNTIF('Simulation Results'!{sc}2:{sc}{sim_end},A{i})"
        summ[f"C{i}"] = f"=COUNTIF('Manual Test Cases'!{mc}2:{mc}{len(man) + 60},A{i})"
    summ.column_dimensions["A"].width = 40
    summ.column_dimensions["B"].width = 34
    summ.column_dimensions["C"].width = 12
    if cats:
        ch = BarChart()
        ch.type, ch.title, ch.height, ch.width = "bar", "Cases per category", 9, 16
        ch.add_data(Reference(summ, min_col=2, max_col=3, min_row=14, max_row=14 + len(cats)),
                    titles_from_data=True)
        ch.set_categories(Reference(summ, min_col=1, min_row=15, max_row=14 + len(cats)))
        summ.add_chart(ch, "E2")
    wb.move_sheet("Summary", offset=-wb.index(summ))
    wb.save(out)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2])
